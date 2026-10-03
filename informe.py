"""
Informe: Excel con varias hojas + gráficos PNG.
"""
from __future__ import annotations

import textwrap
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from matplotlib.colors import LinearSegmentedColormap  # noqa: E402
from scipy.cluster.hierarchy import dendrogram  # noqa: E402

from optimizacion import frontera_eficiente, pesos_max_sharpe, pesos_min_var  # noqa: E402

# Paleta categórica validada (orden fijo) y tinta de texto
SERIES = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948"]
TINTA, TINTA_2, GRIS, FONDO = "#0b0b0b", "#52514e", "#d9d8d4", "#fcfcfb"

plt.rcParams.update({
    "figure.facecolor": FONDO, "axes.facecolor": FONDO, "savefig.facecolor": FONDO,
    "axes.edgecolor": GRIS, "axes.labelcolor": TINTA_2, "xtick.color": TINTA_2,
    "ytick.color": TINTA_2, "text.color": TINTA, "axes.grid": True, "grid.color": "#ecebe8",
    "grid.linewidth": 0.8, "axes.spines.top": False, "axes.spines.right": False,
    "axes.axisbelow": True, "font.size": 10, "axes.titlesize": 12, "axes.titleweight": "bold", "axes.titlelocation": "left",
})


def _corto(texto: str, n: int = 38) -> str:
    texto = str(texto)
    return texto if len(texto) <= n else texto[: n - 1] + "…"


# ----------------------------------------------------------------------------
# Gráficos
# ----------------------------------------------------------------------------
def grafico_pesos(cartera: pd.DataFrame, ruta: Path) -> Path:
    df = cartera.sort_values("peso", ascending=True)
    y = np.arange(len(df))
    alto = 0.38
    fig, ax = plt.subplots(figsize=(9, 0.45 * len(df) + 1.6))
    ax.barh(y + alto / 2, df["peso"] * 100, height=alto, color=SERIES[0], label="Peso")
    ax.barh(y - alto / 2, df["contrib_riesgo"] * 100, height=alto, color=SERIES[1],
            label="Contribución al riesgo")
    ax.set_yticks(y, [_corto(n) for n in df["nombre"]])
    ax.set_xlabel("%")
    ax.grid(axis="y", visible=False)
    for yi, v in zip(y, df["peso"] * 100):
        ax.text(v + 0.3, yi + alto / 2, f"{v:.1f}", va="center", fontsize=8, color=TINTA_2)
    ax.legend(loc="lower right", frameon=False)
    ax.set_title("Cartera: peso y contribución al riesgo por fondo")
    fig.tight_layout()
    fig.savefig(ruta, dpi=150)
    plt.close(fig)
    return ruta


def grafico_dendrograma(enlace, etiquetas: list[str], ruta: Path) -> Path:
    fig, ax = plt.subplots(figsize=(9, 0.35 * len(etiquetas) + 1.5))
    dendrogram(enlace, labels=[_corto(e, 45) for e in etiquetas], orientation="right", ax=ax,
               color_threshold=0, above_threshold_color=SERIES[0])
    ax.set_xlabel("Distancia de enlace (a partir de la distancia de correlación √(½(1−ρ)))")
    ax.grid(axis="y", visible=False)
    ax.set_title("Estructura jerárquica de los fondos seleccionados")
    fig.tight_layout()
    fig.savefig(ruta, dpi=150)
    plt.close(fig)
    return ruta


def grafico_riesgo_rentabilidad(tabla: pd.DataFrame, seleccion: list[str], ruta: Path) -> Path:
    cand = tabla[tabla["vol_ajustada"].notna() & tabla["rent_anual"].notna()
                & (tabla["estado"] != "excluido: datos anómalos")]
    otros = cand.drop(index=[s for s in seleccion if s in cand.index])
    sel = cand.loc[[s for s in seleccion if s in cand.index]]
    fig, ax = plt.subplots(figsize=(9, 6))
    ax.scatter(otros["vol_ajustada"] * 100, otros["rent_anual"] * 100, s=14, color="#b9b8b3",
               alpha=0.6, linewidths=0, label="Resto del universo")
    ax.scatter(sel["vol_ajustada"] * 100, sel["rent_anual"] * 100, s=60, color=SERIES[0],
               edgecolors=FONDO, linewidths=2, label="Seleccionados", zorder=3)
    # etiquetas sin solaparse: desplazamiento vertical greedy en coordenadas de pantalla
    fig.canvas.draw()
    colocadas = []
    for isin, f in sel.sort_values("rent_anual", ascending=False).iterrows():
        x, y = f["vol_ajustada"] * 100, f["rent_anual"] * 100
        px, py = ax.transData.transform((x, y))
        dy = 4
        while any(abs(px - qx) < 120 and abs(py + dy - qy) < 11 for qx, qy in colocadas):
            dy -= 11
        colocadas.append((px, py + dy))
        ax.annotate(_corto(f.get("nombre", isin), 24), (x, y), xytext=(6, dy), textcoords="offset points",
                    fontsize=7.5, color=TINTA_2,
                    arrowprops=dict(arrowstyle="-", color=GRIS, lw=0.6) if dy < -4 else None)
    ax.set_xlabel("Volatilidad anual ajustada por autocorrelación (%)")
    ax.set_ylabel("Rentabilidad anual (%)")
    ax.legend(frameon=False, loc="upper left")
    ax.set_title("Riesgo y rentabilidad en la ventana de estimación")
    fig.tight_layout()
    fig.savefig(ruta, dpi=150)
    plt.close(fig)
    return ruta


def grafico_correlaciones(corr: pd.DataFrame, nombres: pd.Series, ruta: Path) -> Path:
    cmap = LinearSegmentedColormap.from_list("correlacion", [SERIES[0], FONDO, SERIES[1]])
    etiquetas = [_corto(nombres.get(c, c), 28) for c in corr.columns]
    n = len(corr)
    fig, ax = plt.subplots(figsize=(0.55 * n + 2.2, 0.55 * n + 2))
    im = ax.imshow(corr.values, cmap=cmap, vmin=-1, vmax=1)
    ax.set_xticks(range(n), etiquetas, rotation=45, ha="right", fontsize=8)
    ax.set_yticks(range(n), etiquetas, fontsize=8)
    for i in range(n):
        for j in range(n):
            v = corr.values[i, j]
            color = FONDO if abs(v) > 0.55 else TINTA
            ax.text(j, i, f"{v:.2f}", ha="center", va="center", fontsize=7, color=color)
    ax.set_xticks(np.arange(-0.5, n, 1), minor=True)
    ax.set_yticks(np.arange(-0.5, n, 1), minor=True)
    ax.grid(which="minor", color=FONDO, linewidth=1.5)
    ax.grid(which="major", visible=False)
    for spine in ax.spines.values():
        spine.set_visible(False)
    cbar = fig.colorbar(im, ax=ax, shrink=0.8)
    cbar.ax.tick_params(labelsize=8)
    ax.set_title("Matriz de correlaciones entre las posiciones")
    fig.tight_layout()
    fig.savefig(ruta, dpi=150)
    plt.close(fig)
    return ruta


def grafico_frontera_eficiente(rend: pd.DataFrame, cov: pd.DataFrame, w: pd.Series, periodos: int,
                               rf_periodo: float, nombres: pd.Series, ruta: Path) -> Path:
    mu = rend.mean() * periodos
    cov_an = cov * periodos
    frontera = frontera_eficiente(cov_an, mu)

    def vol_ret(pesos: pd.Series) -> tuple[float, float]:
        p = pesos.reindex(cov_an.index).fillna(0).values
        return float(np.sqrt(p @ cov_an.values @ p)), float(p @ mu.values)

    vol_cartera, ret_cartera = vol_ret(w)
    vol_gmv, ret_gmv = vol_ret(pesos_min_var(cov_an, 1.0))
    exceso = mu - rf_periodo * periodos
    vol_tan, ret_tan = vol_ret(pesos_max_sharpe(cov_an, exceso, 1.0, shrink_media=0.0))

    vol_fondos = pd.Series({c: np.sqrt(cov_an.loc[c, c]) for c in cov_an.index}) * 100
    ret_fondos = mu.reindex(cov_an.index) * 100

    fig, ax = plt.subplots(figsize=(9, 6))
    ax.plot(frontera["vol"] * 100, frontera["ret"] * 100, color=SERIES[0], lw=2,
            label="Frontera eficiente (fondos seleccionados)")
    ax.scatter(vol_fondos, ret_fondos, s=40, color="#b9b8b3", zorder=2,
               label="Fondos seleccionados (individual)")
    # etiquetas sin solaparse: mismo desplazamiento vertical greedy que grafico_riesgo_rentabilidad
    fig.canvas.draw()
    colocadas = []
    for isin in ret_fondos.sort_values(ascending=False).index:
        x, y = vol_fondos[isin], ret_fondos[isin]
        px, py = ax.transData.transform((x, y))
        dy = 4
        while any(abs(px - qx) < 130 and abs(py + dy - qy) < 11 for qx, qy in colocadas):
            dy -= 11
        colocadas.append((px, py + dy))
        ax.annotate(_corto(nombres.get(isin, isin), 22), (x, y), xytext=(6, dy), textcoords="offset points",
                    fontsize=7, color=TINTA_2,
                    arrowprops=dict(arrowstyle="-", color=GRIS, lw=0.6) if dy < -4 else None)
    ax.scatter([vol_gmv * 100], [ret_gmv * 100], s=90, marker="D", color=SERIES[2],
               zorder=3, label="Mínima varianza")
    ax.scatter([vol_tan * 100], [ret_tan * 100], s=90, marker="^", color=SERIES[3],
               zorder=3, label="Máximo Sharpe (tangente)")
    ax.scatter([vol_cartera * 100], [ret_cartera * 100], s=220, marker="*", color=SERIES[1],
               edgecolors=TINTA, linewidths=0.6, zorder=4, label="Cartera (HRP, con límites 3-20 %)")
    ax.set_xlabel("Volatilidad anual (%)")
    ax.set_ylabel("Rentabilidad esperada anual (%)")
    ax.legend(frameon=False, loc="upper left", fontsize=8.5)
    ax.set_title("Frontera eficiente de las posiciones seleccionadas")
    fig.tight_layout()
    fig.savefig(ruta, dpi=150)
    plt.close(fig)
    return ruta


def grafico_backtest(serie: pd.DataFrame, ruta: Path) -> Path:
    riqueza = (1 + serie.fillna(0)).cumprod() * 100
    fig, ax = plt.subplots(figsize=(9, 5))
    nombres = {"hrp": "HRP", "inversa_vol": "Inversa de la volatilidad", "min_var": "Mínima varianza",
               "max_sharpe": "Máximo Sharpe", "equiponderada": "1/N (seleccionados)",
               "universo_1N": "1/N (todo el universo)"}
    for i, c in enumerate(riqueza.columns):
        ax.plot(riqueza.index, riqueza[c], lw=2, color=SERIES[i % len(SERIES)],
                label=nombres.get(c, c), ls="--" if c == "universo_1N" else "-")
    ax.set_ylabel("Valor (base 100)")
    ax.legend(frameon=False, loc="upper left", fontsize=9)
    ax.set_title("Backtest fuera de muestra (walk-forward)")
    fig.tight_layout()
    fig.savefig(ruta, dpi=150)
    plt.close(fig)
    return ruta


# ----------------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------------
FORMATO_PCT = {"peso", "contrib_riesgo", "rent_anual", "vol_anual", "vol_ajustada", "max_drawdown",
               "cobertura", "ter", "psr", "rotacion_media"}


def exportar_excel(ruta: Path, hojas: dict[str, pd.DataFrame], imagenes: list[Path]) -> Path:
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        for nombre, df in hojas.items():
            df.to_excel(xw, sheet_name=nombre[:31])
        libro = xw.book
        cabecera = PatternFill("solid", fgColor="E8EEF7")
        for nombre, df in hojas.items():
            ws = libro[nombre[:31]]
            ws.freeze_panes = "B2"
            columnas = [df.index.name or ""] + list(df.columns)
            for j, col in enumerate(columnas, start=1):
                celda = ws.cell(row=1, column=j)
                celda.font = Font(bold=True)
                celda.fill = cabecera
                celda.alignment = Alignment(wrap_text=True, vertical="center")
                ancho = max([len(str(col))] + [len(str(v)) for v in (df.index if j == 1 else df.iloc[:, j - 2]).astype(str)[:200]])
                ws.column_dimensions[get_column_letter(j)].width = min(max(10, ancho + 2), 55)
                if str(col) in FORMATO_PCT or str(col).startswith("estilo_") and col != "estilo_R2":
                    for fila in range(2, len(df) + 2):
                        ws.cell(row=fila, column=j).number_format = "0.0%"
                elif j > 1 and pd.api.types.is_float_dtype(df.iloc[:, j - 2]):
                    for fila in range(2, len(df) + 2):
                        ws.cell(row=fila, column=j).number_format = "0.00"
        if imagenes:
            ws = libro.create_sheet("Gráficos")
            fila = 1
            for img in imagenes:
                im = XLImage(str(img))
                escala = 780 / im.width
                im.width, im.height = im.width * escala, im.height * escala
                ws.add_image(im, f"A{fila}")
                fila += int(im.height / 20) + 3
    return ruta


def configuracion_plana(cfg: dict, prefijo: str = "") -> pd.DataFrame:
    filas = []
    for k, v in cfg.items():
        clave = f"{prefijo}{k}"
        if isinstance(v, dict):
            filas.append(configuracion_plana(v, clave + ".").reset_index())
        else:
            filas.append(pd.DataFrame({"parametro": [clave], "valor": [str(v)]}))
    df = pd.concat(filas, ignore_index=True) if filas else pd.DataFrame(columns=["parametro", "valor"])
    if "index" in df.columns:
        df = df.drop(columns="index")
    return df.set_index("parametro")


def texto_resumen(cartera: pd.DataFrame, stats: dict, exposicion: pd.Series) -> str:
    lineas = [
        f"Fondos en cartera: {len(cartera)}",
        f"Volatilidad esperada (covarianza estimada): {stats['vol_esperada']:.1%}",
        f"Histórico de la cartera en la ventana — rentabilidad {stats['rent_hist']:.1%}, "
        f"volatilidad {stats['vol_hist']:.1%}, Sharpe {stats['sharpe_hist']:.2f}, "
        f"caída máxima {stats['dd_hist']:.1%}",
        "Exposición por tipo de activo: " + ", ".join(f"{k} {v:.0%}" for k, v in exposicion.items()),
    ]
    return "\n".join(textwrap.fill(l, 100) for l in lineas)
